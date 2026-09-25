from openpyxl import load_workbook
from ..models import Word, WordStats
from ..extensions import db


def read_excel(filepath):
    wb = load_workbook(filepath, read_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        return []

    headers = [str(h).strip().lower() if h else "" for h in rows[0]]
    words = []

    for row in rows[1:]:
        row_dict = dict(zip(headers, row))
        eng = row_dict.get("english") or row_dict.get("inglés") or row_dict.get("ingles")
        esp = row_dict.get("spanish") or row_dict.get("español") or row_dict.get("espanol")

        if not eng or not esp:
            continue

        eng = str(eng).strip()
        esp = str(esp).strip()

        notes_cell = (
            row_dict.get("notes")
            or row_dict.get("observación")
            or row_dict.get("observacion")
            or row_dict.get("pista")
            or row_dict.get("hint")
        )
        notes = str(notes_cell).strip() if notes_cell is not None and str(notes_cell).strip() else None

        if eng and esp:
            words.append({"english": eng, "spanish": esp, "notes": notes})

    return words


def import_words(words):
    added = 0
    skipped = 0

    for w in words:
        exists = Word.query.filter_by(english=w["english"]).first()
        if exists:
            skipped += 1
            continue

        word = Word(english=w["english"], spanish=w["spanish"], notes=w.get("notes"))
        db.session.add(word)
        db.session.flush()

        stats = WordStats(word_id=word.id)
        db.session.add(stats)
        added += 1

    db.session.commit()
    return added, skipped


def import_file(filepath):
    words = read_excel(filepath)
    return import_words(words)
