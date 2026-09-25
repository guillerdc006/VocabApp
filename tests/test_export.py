import io

from openpyxl import load_workbook

from app.models import Word


class TestExport:
    def test_exporta_archivo_con_datos(self, client, make_word):
        make_word("apple", "manzana", notes="fruta", practiced=4, correct=3)
        make_word("dog", "perro")

        resp = client.get("/vocabulary/export")
        assert resp.status_code == 200
        assert "attachment" in resp.headers.get("Content-Disposition", "")
        assert "vocabulario_export.xlsx" in resp.headers.get("Content-Disposition", "")

        wb = load_workbook(io.BytesIO(resp.data))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        header = [str(c) for c in rows[0]]

        assert "english" in header
        assert "spanish" in header
        assert "notes" in header
        assert "times_practiced" in header
        assert "times_correct" in header
        assert "accuracy_percent" in header
        assert "next_review" in header

        data = {r[0]: r for r in rows[1:]}
        row_apple = data["apple"]
        assert row_apple[1] == "manzana"
        assert row_apple[2] == "fruta"
        assert row_apple[5] == 4
        assert row_apple[6] == 3
        assert row_apple[7] == 75

        row_dog = data["dog"]
        assert row_dog[2] == "" or row_dog[2] is None
        assert row_dog[5] == 0

    def test_export_roundtrip_sin_duplicados(self, client, make_word):
        make_word("apple", "manzana", notes="fruta")
        export = client.get("/vocabulary/export").data

        resp = client.post(
            "/vocabulary/import",
            data={"file": (io.BytesIO(export), "vocabulario_export.xlsx")},
            content_type="multipart/form-data",
        )
        assert resp.status_code == 302

        with client.application.app_context():
            assert Word.query.count() == 1
            word = Word.query.filter_by(english="apple").first()
            assert word.spanish == "manzana"
            assert word.notes == "fruta"